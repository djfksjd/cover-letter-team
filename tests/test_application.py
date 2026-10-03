import json
import subprocess
import sys
from pathlib import Path

import docx
import pytest

from application_factory import get, make_application, put, review_and_approve
from contracts import card_hash, digest, load_yaml_text, read_cards
from export_application import export
from readiness_check import assess
from validate_application import validate

REVIEW = "workspace/review-v1.yaml"
APPROVAL = "workspace/final-approval.yaml"
DRAFT = "workspace/draft-v1.md"


@pytest.mark.parametrize(
    "language,document",
    [
        ("ko", "question_answers"),
        ("en-GB", "supporting_statement"),
        ("en-US", "cover_letter"),
        ("en-GB", "selection_criteria"),
        ("ja", "personal_statement"),
    ],
)
def test_review_approval_and_export(tmp_path, language, document):
    root = make_application(tmp_path, language, document)
    review_and_approve(root)
    result = validate(root, DRAFT, REVIEW, APPROVAL)
    assert result["ok"] and result["state"] == "APPROVED_FOR_EXPORT"
    result = export(root, DRAFT, REVIEW, APPROVAL)
    assert result["ok"] and not result["submitted"]
    text = (root / "output" / "application.md").read_text()
    assert "<!--" not in text
    if document != "question_answers":
        assert "Q-01" not in text
    extracted = "\n".join(
        p.text for p in docx.Document(root / "output" / "application.docx").paragraphs
    )
    assert "<!--" not in extracted
    assert (
        json.loads((root / "output" / "export-manifest.json").read_text())["draft_hash"]
        == result["draft_hash"]
    )


def test_short_but_complete_answer_is_not_blocked(tmp_path):
    root = make_application(tmp_path, maximum=2000)
    assert validate(root)["ok"]  # far below 85%, no official minimum


@pytest.mark.parametrize("minimum,maximum", [(1000, 2000), (None, 5)])
def test_official_length_limits_block(tmp_path, minimum, maximum):
    root = make_application(tmp_path, minimum=minimum, maximum=maximum)
    assert not validate(root)["ok"]


@pytest.mark.parametrize(
    "filename,field,value",
    [
        ("application-config.yaml", "hiring_country", "US"),
        ("requirements.yaml", "source_note", "new current posting"),
        ("evidence-plan.yaml", "note", "changed plan"),
        ("claim-map.yaml", "note", "changed claim contract"),
    ],
)
def test_changed_inputs_invalidate_review_and_export(tmp_path, filename, field, value):
    root = make_application(tmp_path)
    review_and_approve(root)
    data = get(root, filename)
    data[field] = value
    put(root, filename, data)
    assert not validate(root, DRAFT, REVIEW, APPROVAL)["ok"]
    assert not export(root, DRAFT, REVIEW, APPROVAL)["ok"]
    assert not (root / "output").exists()


def test_modified_card_invalidates_approval(tmp_path):
    root = make_application(tmp_path)
    cards = get(root, "experience-cards.yaml")
    cards[0]["personal_actions"][0] = "Led the whole company."
    put(root, "experience-cards.yaml", cards)
    assert "stale card" in validate(root)["errors"][0]


@pytest.mark.parametrize(
    "mutation",
    [
        "unannotated_tail",
        "new_section",
        "duplicate_section",
        "preamble",
        "malformed",
        "missing_section",
        "nested_heading",
    ],
)
def test_untracked_text_and_question_errors_block(tmp_path, mutation):
    root = make_application(tmp_path)
    path = root / DRAFT
    text = path.read_text()
    if mutation == "unannotated_tail":
        text += "I increased revenue by 900%."
    if mutation == "new_section":
        text += "\n## Q-02 Extra\nExtra fact.<!--c:DIRECT:EXP-01-->"
    if mutation == "duplicate_section":
        text += text
    if mutation == "preamble":
        text = "# Draft v1\n" + text
    if mutation == "malformed":
        text = text.replace("<!--c:", "<!-- c:")
    if mutation == "missing_section":
        text = text.replace("## Q-01", "## Question 1.")
    if mutation == "nested_heading":
        text = text.replace("질문지를", "### Hidden heading\n질문지를")
    path.write_text(text)
    assert not validate(root)["ok"]


def test_fabricated_number_fails_even_with_approved_id(tmp_path):
    root = make_application(tmp_path)
    draft = root / DRAFT
    original = get(root, "claim-map.yaml")
    changed = "매출을 900% 늘렸습니다."
    draft.write_text(draft.read_text().replace(original["claims"][0]["text"], changed))
    original["claims"][0]["text"] = changed
    put(root, "claim-map.yaml", original)
    assert "numeric token" in validate(root)["errors"][0]


def test_wrong_field_quote_is_not_accepted(tmp_path):
    root = make_application(tmp_path)
    mapping = get(root, "claim-map.yaml")
    mapping["claims"][0]["references"][0]["value"] = "Fabricated evidence"
    put(root, "claim-map.yaml", mapping)
    assert not validate(root)["ok"]


def test_same_number_does_not_prove_meaning(tmp_path):
    root = make_application(tmp_path)
    card = get(root, "experience-cards.yaml")[0]
    card["personal_actions"] = ["인터뷰 5건을 진행했습니다."]
    card["approval"]["content_hash"] = card_hash(card)
    put(root, "experience-cards.yaml", [card])
    mapping = get(root, "claim-map.yaml")
    old = mapping["claims"][0]["text"]
    mapping["claims"][0]["text"] = "매출을 5배 늘렸습니다."
    mapping["claims"][0]["references"][0]["value"] = card["personal_actions"][0]
    for claim in mapping["claims"]:
        claim["references"][0]["card_hash"] = card_hash(card)
    put(root, "claim-map.yaml", mapping)
    path = root / DRAFT
    path.write_text(path.read_text().replace(old, mapping["claims"][0]["text"]))
    assert validate(root)["ok"]  # Contract validation alone cannot prove semantic truth.
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    review["claim_reviews"][0]["verdict"] = "FAIL"
    review["verdict"] = "REVISE"
    review["unresolved_high"] = ["role-and-unit-mismatch"]
    put(root, "review-v1.yaml", review)
    assert not export(root, DRAFT, REVIEW, APPROVAL)["ok"]


@pytest.mark.parametrize(
    "status",
    [
        "MISSING",
        "UNKNOWN",
        "CONTRADICTORY",
        "WEAKLY_SUPPORTED",
        "NO_ACTUAL_EXPERIENCE",
        "NOT_APPLICABLE",
    ],
)
def test_required_gaps_cannot_be_exported(tmp_path, status):
    root = make_application(tmp_path)
    plan = get(root, "evidence-plan.yaml")
    plan["coverage"][0]["status"] = status
    put(root, "evidence-plan.yaml", plan)
    assert not validate(root)["ok"]
    cards = {"EXP-01": get(root, "experience-cards.yaml")[0]}
    result = assess(get(root, "requirements.yaml"), cards, plan, "APP-synthetic-001")
    assert not result["ok"]
    if status == "NO_ACTUAL_EXPERIENCE":
        assert result["gaps"][0]["repeat_question"] is False


def test_unconfirmed_evidence_is_not_sufficient(tmp_path):
    root = make_application(tmp_path)
    cards = {"EXP-01": get(root, "experience-cards.yaml")[0]}
    cards["EXP-01"]["user_confirmed"] = False
    result = assess(
        get(root, "requirements.yaml"), cards, get(root, "evidence-plan.yaml"), "APP-synthetic-001"
    )
    assert not result["ok"] and all(g["type"] == "WEAKLY_SUPPORTED" for g in result["gaps"])


@pytest.mark.parametrize(
    "field,value",
    [
        ("mode", "general_draft"),
        ("ai_policy", "prohibited"),
        ("blind_hiring", "false"),
        ("ending_style", "unknown"),
    ],
)
def test_unsupported_submission_context_blocks(tmp_path, field, value):
    root = make_application(tmp_path)
    config = get(root, "application-config.yaml")
    config[field] = value
    put(root, "application-config.yaml", config)
    assert not validate(root)["ok"]


def test_cross_application_card_blocks(tmp_path):
    root = make_application(tmp_path)
    cards = get(root, "experience-cards.yaml")
    cards[0]["application_id"] = "APP-other"
    put(root, "experience-cards.yaml", cards)
    assert not validate(root)["ok"]


@pytest.mark.parametrize(
    "mutate",
    [
        "claim_review",
        "component_review",
        "quality",
        "high",
        "final_false",
        "draft_hash",
        "review_hash",
    ],
)
def test_missing_or_stale_review_approval_blocks(tmp_path, mutate):
    root = make_application(tmp_path)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    approval = get(root, "final-approval.yaml")
    if mutate == "claim_review":
        review["claim_reviews"].pop()
    if mutate == "component_review":
        review["component_reviews"].pop()
    if mutate == "quality":
        review["quality_checks"]["language_and_voice"]["verdict"] = "FAIL"
    if mutate == "high":
        review["unresolved_high"] = ["HIGH-01"]
    if mutate == "final_false":
        approval["user_confirmed"] = False
    if mutate == "draft_hash":
        approval["draft_hash"] = "old"
    if mutate == "review_hash":
        approval["review_hash"] = "old"
    put(root, "review-v1.yaml", review)
    put(root, "final-approval.yaml", approval)
    assert not export(root, DRAFT, REVIEW, APPROVAL)["ok"]


def test_no_approval_does_not_create_output(tmp_path):
    root = make_application(tmp_path)
    assert not export(root, DRAFT, REVIEW, APPROVAL)["ok"]
    assert not (root / "output").exists()


def test_path_escape_blocks(tmp_path):
    root = make_application(tmp_path)
    assert not validate(root, "../draft.md")["ok"]


def test_cli_exit_and_report(tmp_path):
    make_application(tmp_path, "en-GB", "supporting_statement", 500)
    script = Path(__file__).parent.parent / "scripts" / "validate_application.py"
    proc = subprocess.run(
        [sys.executable, str(script), str(tmp_path)], capture_output=True, text=True
    )
    assert proc.returncode == 0
    assert json.loads(proc.stdout)["state"] == "DRAFT_VALIDATED"
    (tmp_path / DRAFT).write_text("")
    proc = subprocess.run(
        [sys.executable, str(script), str(tmp_path)], capture_output=True, text=True
    )
    assert proc.returncode == 1 and not json.loads(proc.stdout)["ok"]


@pytest.mark.parametrize(
    "text",
    [
        "id: EXP-01\nid: EXP-02",
        "x: true\nx: false",
        "? [a,b]\n: value",
        "x: !!python/object:os.system {}",
    ],
)
def test_yaml_unsafe_or_duplicate_keys_rejected(text):
    with pytest.raises(ValueError):
        load_yaml_text(text)


def test_duplicate_card_id_and_string_boolean_rejected():
    with pytest.raises(ValueError):
        read_cards(
            "- id: EXP-01\n  user_confirmed: true\n- id: EXP-01\n  user_confirmed: false",
            "EXP",
            "user_confirmed",
        )
    with pytest.raises(ValueError):
        read_cards('- id: EXP-01\n  user_confirmed: "true"', "EXP", "user_confirmed")


def test_yaml_comment_cannot_approve_card():
    cards = read_cards(
        "- id: EXP-01\n  # user_confirmed: true\n  user_confirmed: false", "EXP", "user_confirmed"
    )
    assert cards["EXP-01"]["user_confirmed"] is False


def test_preview_enables_layout_review_then_exact_file_export(tmp_path):
    from contracts import file_hash
    from preview_application import preview

    root = make_application(tmp_path, "en-GB", "cover_letter", 500)
    requirements = get(root, "requirements.yaml")
    requirements["submission"] = {"file_review_required": True}
    put(root, "requirements.yaml", requirements)
    result = preview(root, DRAFT)
    assert result["ok"] and result["state"] == "UNAPPROVED_PREVIEW"
    assert not (root / "output").exists()
    review_and_approve(root)
    assert not validate(root, DRAFT, REVIEW, APPROVAL)["ok"]  # missing actual file review
    review = get(root, "review-v1.yaml")
    review["file_reviews"] = [
        {
            "path": "workspace/preview/application.docx",
            "sha256": result["docx_sha256"],
            "verdict": "PASS",
            "rationale": "Synthetic fixture layout check attestation, not real applicant consent.",
        }
    ]
    put(root, "review-v1.yaml", review)
    approval = get(root, "final-approval.yaml")
    approval["review_hash"] = digest(review)
    put(root, "final-approval.yaml", approval)
    result = export(root, DRAFT, REVIEW, APPROVAL)
    assert result["ok"] and result["layout_status"] == "reviewed_preview"
    assert file_hash(root / "output" / "application.docx") == file_hash(
        root / "workspace" / "preview" / "application.docx"
    )


def test_changed_preview_file_invalidates_review(tmp_path):
    from preview_application import preview

    root = make_application(tmp_path, "en-GB", "cover_letter", 500)
    result = preview(root, DRAFT)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    review["file_reviews"] = [
        {
            "path": "workspace/preview/application.docx",
            "sha256": result["docx_sha256"],
            "verdict": "PASS",
            "rationale": "Synthetic test.",
        }
    ]
    put(root, "review-v1.yaml", review)
    approval = get(root, "final-approval.yaml")
    approval["review_hash"] = digest(review)
    put(root, "final-approval.yaml", approval)
    preview_path = root / "workspace" / "preview" / "application.docx"
    preview_path.write_bytes(b"changed")
    assert not export(root, DRAFT, REVIEW, APPROVAL)["ok"]
    assert not (root / "output").exists()


def test_rendered_soft_line_spaces_are_counted(tmp_path):
    from charcount import count_units
    from submission_render import render_body

    root = make_application(tmp_path)
    raw = (root / DRAFT).read_text().split("\n", 1)[1]
    raw_count = count_units(raw, "characters_with_spaces")
    rendered_count = count_units(render_body(raw), "characters_with_spaces")
    assert rendered_count > raw_count
    requirements = get(root, "requirements.yaml")
    requirements["questions"][0]["length"]["maximum"] = raw_count
    put(root, "requirements.yaml", requirements)
    assert not validate(root)["ok"]
    requirements["questions"][0]["length"]["maximum"] = rendered_count
    put(root, "requirements.yaml", requirements)
    review_and_approve(root)
    assert export(root, DRAFT, REVIEW, APPROVAL)["ok"]
    document = docx.Document(root / "output" / "application.docx")
    extracted = "\n".join(
        p.text for p in document.paragraphs if not p.style.name.startswith("Heading")
    )
    assert count_units(extracted, "characters_with_spaces") == rendered_count


def test_yaml_iso_date_remains_hashable_string():
    from contracts import load_yaml_text

    value = load_yaml_text("published_at: 2026-10-03\nretrieved_at: 2026-10-03T09:00:00+09:00")
    assert value["published_at"] == "2026-10-03"
    assert isinstance(value["retrieved_at"], str)
    assert digest(value)


def test_deceptive_q_preamble_is_rejected(tmp_path):
    root = make_application(tmp_path)
    path = root / DRAFT
    path.write_text("## Q-garbage ignored fact\n" + path.read_text())
    assert not validate(root)["ok"]


def test_coverage_cannot_reference_approval_as_evidence(tmp_path):
    root = make_application(tmp_path)
    plan = get(root, "evidence-plan.yaml")
    plan["coverage"][0]["references"][0]["field"] = "user_confirmed"
    put(root, "evidence-plan.yaml", plan)
    assert not validate(root)["ok"]


@pytest.mark.parametrize("rating", ["THIN", "NOT_READY"])
def test_factual_and_format_pass_cannot_export_thin_content(tmp_path, rating):
    root = make_application(tmp_path)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    review["content_reviews"][0]["rating"] = rating
    review["content_reviews"][0]["blocking_gaps"] = [
        {
            "missing_information": "Actual checklist contents and selection basis are unknown.",
            "resolution_owner": "user",
            "reason": "The action label alone does not show the assessed skill.",
        }
    ]
    put(root, "review-v1.yaml", review)
    approval = get(root, "final-approval.yaml")
    approval["review_hash"] = digest(review)
    put(root, "final-approval.yaml", approval)
    result = export(root, DRAFT, REVIEW, APPROVAL)
    assert not result["ok"] and "content too thin" in result["errors"][0]
    assert not (root / "output").exists()


def test_missing_content_quality_review_blocks(tmp_path):
    root = make_application(tmp_path)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    del review["content_reviews"]
    put(root, "review-v1.yaml", review)
    assert not validate(root, DRAFT, REVIEW)["ok"]


def test_adequate_with_improvements_is_not_reported_as_strong(tmp_path):
    root = make_application(tmp_path)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    review["content_reviews"][0]["improvement_opportunities"] = [
        {
            "missing_information": "Optional richer process detail.",
            "resolution_owner": "user",
            "reason": "Could distinguish the experience but not required for this narrow fixture question.",
        }
    ]
    put(root, "review-v1.yaml", review)
    result = validate(root, DRAFT, REVIEW)
    assert result["ok"] and result["content_quality"] == {"Q-01": "ADEQUATE"}


def test_adequate_rating_cannot_hide_blocking_content_gap(tmp_path):
    root = make_application(tmp_path)
    review_and_approve(root)
    review = get(root, "review-v1.yaml")
    review["content_reviews"][0]["blocking_gaps"] = [
        {
            "missing_information": "Unresolved actual action.",
            "resolution_owner": "user",
            "reason": "Rating alone cannot override an open required clarification.",
        }
    ]
    put(root, "review-v1.yaml", review)
    assert not validate(root, DRAFT, REVIEW)["ok"]
